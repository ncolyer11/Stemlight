# made by ncolyer on the 17/03/2023

# A program that takes an input of a 7x27x7 (XYZ) schematic in the .litematic format, representing the blocks a given
# nether tree farm layout harvests per cycle, and decodes the 3D data inside that file to compare it to pre-calculated
# heatmaps in a .xlsx file to work out the average stems, shroomlights and wart blocks harvested per cycle for that
# nether tree farm layout

# for more info on how this & the heatmaps were made visit the huge fungi huggers discord: https://discord.gg/EKKkyfcPPV

# if the below import functions are underlined in red run `pip install litemapy` and `pip install openpyxl`
from litemapy import Schematic
import openpyxl
import time

print("\n"
      "\t\t╔════════════════════════════════════════════════════╗\n"
      "\t\t  NOTE: Make sure your litematic selection measures\n"
      "\t\t  7x27x7 and the base of the selection is centred on\n"
      "\t\t        the block directly above the nylium\n"
      "\t\t╚════════════════════════════════════════════════════╝\n")

# taking file input from user and checking to see if the path and schematic size is valid
while True:
    user_name = input(r"Enter your Windows username (e.g. for C:\Users\Dylan\ 'Dylan' would be what you enter): ")
    schematic_name = input("Enter your layout schematic name (excluding the '.litematic): ")
    try:
        schem = Schematic.load(f"C:\\Users\\{user_name}\\AppData\\Roaming\\.minecraft\\schematics\\"
                               f"ntf_heatmap_layouts\\{schematic_name}.litematic")
        reg = list(schem.regions.values())[0]
        if len(reg.xrange()) == 7 and len(reg.yrange()) == 27 and len(reg.zrange()) == 7:
            print("\nSchematic loaded succesfully :)\n")
            break
        else:
            print("\nSizeError: Ensure the volume of your schematic measures 7x27x7.\n")
    except FileNotFoundError:
        print("\nNameError: Incorrect user or schematic name, please try again.\n")

start_time = time.time()

workbook = openpyxl.load_workbook('heatmaps.xlsx')
cell_values = {}

# opening the heatmap excel spreadsheet once and writing all the values to an array
for sheet_name in workbook.sheetnames:
    sheet = workbook[sheet_name]
    for row in sheet.iter_rows():
        for cell in row:
            cell_values[(sheet_name, cell.row, cell.column)] = cell.value


# function for accessing a single value/cell in the above array
def get_cell_value(sheet_name, row_number, column_number):
    return cell_values[(sheet_name, row_number, column_number)]


# function that computes the instantaneous value of the avg stems and shroomlights of the layout
def stems_and_shrooms(avg_stems, avg_shroomlights, row, col):
    avg_stems += get_cell_value('stems', row, col)
    avg_shroomlights += get_cell_value('shroomlights', row, col)
    print(f"\t• Stems = {get_cell_value('stems', row, col)}\n"
          f"\t• Shroomlights = {get_cell_value('shroomlights', row, col)}")
    return avg_stems, avg_shroomlights


# function that returns the upper and lower bounds of an array, as an array
def make_bound_array(list):
    return [min(list), max(list)]


# initialising variables
Xrange, Yrange, Zrange = reg.xrange(), reg.yrange(), reg.zrange()
avg_shroomlights, avg_wart_blocks, avg_stems = 0, 0, 0

for y in Yrange:
    for x in Xrange:
        for z in Zrange:
            block = reg.getblock(x, y, z)

            X, Y, Z = x, y, z
            if min(Xrange) < 0:
                X += 6
            if min(Yrange) < 0:
                Y += 26
            if min(Zrange) < 0:
                Z += 6

            if block.blockid != "minecraft:air":
                # 3D data is stored in 'slices' on a 2D spreadsheet, hence some math is needed to convert between them
                col = X + (7 * Z) + 1
                row = 27 - Y
                print(f"At block ({X}, {Y}, {Z}):")
                if block.blockid == "minecraft:red_concrete":
                    avg_stems, avg_shroomlights = stems_and_shrooms(avg_stems, avg_shroomlights, row, col)
                    avg_wart_blocks += get_cell_value('vrm0', row, col)
                    print(f"\t• VRM0 = {get_cell_value('vrm0', row, col)}")
                elif block.blockid == "minecraft:blue_concrete":
                    avg_stems, avg_shroomlights = stems_and_shrooms(avg_stems, avg_shroomlights, row, col)
                    avg_wart_blocks += get_cell_value('vrm1', row, col)
                    print(f"\t• VRM1 = {get_cell_value('vrm1', row, col)}")
                elif block.blockid == "minecraft:cyan_concrete":
                    avg_stems, avg_shroomlights = stems_and_shrooms(avg_stems, avg_shroomlights, row, col)
                    avg_wart_blocks += get_cell_value('vrm2', row, col)
                    print(f"\t• VRM2 = {get_cell_value('vrm2', row, col)}")
                elif block.blockid == "minecraft:light_blue_concrete":
                    avg_stems, avg_shroomlights = stems_and_shrooms(avg_stems, avg_shroomlights, row, col)
                    avg_wart_blocks += get_cell_value('vrm3', row, col)
                    print(f"\t• VRM3 = {get_cell_value('vrm3', row, col)}")
                else:
                    print(f"\t• INVALID blockstate: {block.blockid[10:]}")
stem_E = 24 * avg_stems / 221
shroomlight_E = avg_shroomlights / 2.03192455026454
wart_block_E = avg_wart_blocks / 63.0252319962964 
# wart blocks max doesn't take into account vrm and the actual max warts including VRM still hasn't been found
# well at least the max feasible with vrm as if you can predict the optimal vrm locations for a huge fungi before it
# grows you'd get ~99 wart blocks per fungus, but this isn't possible without rng manipulation

print(f"\n╔═════════════════════════════════════════════════════════════╗\n"
      f"   Avg blocks harvested per fungus:\n"
      f"\t   • Stems: {avg_stems} @{stem_E}E\n"
      f"\t   • Shroomlights: {avg_shroomlights} @{shroomlight_E}E\n"
      f"\t   • Wart Blocks: {avg_wart_blocks} @{wart_block_E}E\n"
      f"   Within the selection: X: {make_bound_array(list(Xrange))} Y: {make_bound_array(list(Yrange))} Z:"
      f" {make_bound_array(list(Zrange))}"
      f"\n╚═════════════════════════════════════════════════════════════╝ ")

end_time = time.time()
elapsed_time = end_time - start_time
print("\nElapsed time:", elapsed_time, "seconds")
